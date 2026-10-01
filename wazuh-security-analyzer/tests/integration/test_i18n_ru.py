"""Full analysis, reports and GUI in Russian."""

import json
import re

import pytest
from openpyxl import load_workbook

from app.core.config import AppConfig
from app.i18n import get_language, set_language
from app.reports.exporters import export_report
from app.services.pipeline import AnalysisPipeline

CYR = re.compile("[а-яА-ЯёЁ]")


@pytest.fixture(scope="module")
def ru_session(tmp_path_factory):
    previous = get_language()
    cfg = AppConfig.model_validate({"ui": {"language": "ru"}})
    session = AnalysisPipeline(cfg, workspace=tmp_path_factory.mktemp("ru_ws"), online=False).run(
        [__import__("pathlib").Path(__file__).resolve().parents[2] / "sample_data"])
    yield session
    session.close()
    set_language(previous)


def test_analysis_texts_are_russian(ru_session):
    set_language("ru")
    s = ru_session.summary
    assert s.executive_summary.startswith("За анализируемый период")
    assert "Возможная компрометация" in s.executive_summary
    groups = ru_session.groups(limit=50)
    for g in groups:
        assert CYR.search(g.what_happened), g.what_happened
        assert CYR.search(g.why_it_matters)
        assert all(CYR.search(e) for e in g.evidence)
        assert all(CYR.search(f.reason) or f.reason.startswith(("CVE", "MITRE")) for f in g.risk_factors)
        for steps in g.recommendations.values():
            assert all(CYR.search(step) for step in steps), steps
    incident = next(i for i in ru_session.incidents() if i.kind == "attack_chain")
    assert incident.title.startswith("Возможная цепочка атаки")
    assert "Доступ к учётным данным" in incident.scenario
    assert incident.assessment == "Potential compromise"  # enum stays English, translated only for display


def test_russian_reports(ru_session, tmp_path):
    set_language("ru")
    written = export_report(ru_session, tmp_path, ["pdf", "html", "docx", "xlsx", "json"], "ru")
    html = (tmp_path / "ru.html").read_text(encoding="utf-8")
    assert '<html lang="ru">' in html and "Резюме для руководства" in html and "Техническое приложение" in html
    assert "КРИТИЧЕСКИЙ" in html
    wb = load_workbook(tmp_path / "ru.xlsx", read_only=True)
    assert "Инциденты" in wb.sheetnames and "Сводка" in wb.sheetnames
    data = json.loads((tmp_path / "ru.json").read_text(encoding="utf-8"))
    assert data["report"]["title"] == "Отчёт об анализе безопасности Wazuh"
    pdf = (tmp_path / "ru.pdf").read_bytes()
    assert pdf.startswith(b"%PDF") and b"DejaVu" in pdf  # Cyrillic-capable embedded font
    assert all(p.exists() for p in written)


@pytest.mark.gui
def test_gui_in_russian(ru_session, tmp_path):
    QtWidgets = pytest.importorskip("PySide6.QtWidgets")
    from app.core.config import ConfigManager
    from app.core.secrets import MemorySecretStore
    from app.database.state import StateStore
    from app.ui.context import AppContext
    from app.ui.main_window import MainWindow
    from app.ui.theme import apply_theme

    set_language("ru")
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    apply_theme(app)
    ctx = AppContext(ConfigManager(tmp_path / "c.yaml"), MemorySecretStore(),
                     StateStore(f"sqlite:///{tmp_path / 's.db'}"))
    window = MainWindow(ctx)
    try:
        assert "Обзор" in window.nav_buttons["dashboard"].text()
        assert "Настройки" in window.nav_buttons["settings"].text()
        assert "MITRE ATT&&CK" in window.nav_buttons["mitre"].text()
        ctx.session = ru_session
        window.pages["alerts"].on_session(ru_session)
        window.navigate("alerts", {"severity": "critical"})
        text = window.pages["alerts"].detail.toPlainText()
        assert "Что произошло?" in text and "Рекомендуемые действия" in text
        model = window.pages["alerts"].findings.model
        assert model.headerData(1, __import__("PySide6.QtCore", fromlist=["Qt"]).Qt.Horizontal) == "Критичность"
    finally:
        ctx.session = None
        window.close()
