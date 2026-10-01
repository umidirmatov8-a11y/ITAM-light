import csv
import json

from docx import Document
from openpyxl import load_workbook

from app.core.config import AppConfig
from app.reports.exporters import export_report
from app.services.pipeline import AnalysisPipeline


def test_all_formats(demo_session, tmp_path):
    written = export_report(demo_session, tmp_path, ["pdf", "html", "docx", "xlsx", "csv", "json"], "r")
    names = {p.name for p in written}
    assert names == {"r.pdf", "r.html", "r.docx", "r.xlsx", "r.csv", "r.json"}
    assert (tmp_path / "r_events.csv").exists()
    assert (tmp_path / "r.pdf").read_bytes().startswith(b"%PDF")
    html = (tmp_path / "r.html").read_text()
    for section in ("Executive Summary", "Statistics", "Incidents", "MITRE ATT&amp;CK", "Recommendations",
                    "Technical Appendix", "Timeline"):
        assert section in html
    data = json.loads((tmp_path / "r.json").read_text())
    assert data["summary"]["events"] == demo_session.summary.events and data["incidents"]
    wb = load_workbook(tmp_path / "r.xlsx", read_only=True)
    assert {"Summary", "Incidents", "Findings", "IOC", "CVE", "MITRE", "Events"} <= set(wb.sheetnames)
    doc = Document(str(tmp_path / "r.docx"))
    assert any("Executive Summary" in p.text for p in doc.paragraphs)
    with (tmp_path / "r_events.csv").open(encoding="utf-8-sig") as fh:
        assert sum(1 for _ in csv.reader(fh)) == demo_session.summary.events + 1


def test_section_selection(demo_session, tmp_path):
    export_report(demo_session, tmp_path, ["html"], "only", sections=("executive_summary",))
    html = (tmp_path / "only.html").read_text()
    assert "Executive Summary" in html and "Technical Appendix" not in html


def test_injection_is_neutralised(tmp_path, jsonl_writer, alert_factory):
    evil = jsonl_writer(tmp_path / "evil.jsonl", [alert_factory(
        rule_id="100999", level=12, description="=HYPERLINK(\"http://x\")<script>alert(1)</script>",
        agent="<img src=x onerror=alert(1)>", data={"srcip": "203.0.113.200"})])
    session = AnalysisPipeline(AppConfig(), workspace=tmp_path / "ws").run([evil])
    export_report(session, tmp_path, ["html", "csv", "xlsx", "pdf", "docx"], "e")
    html = (tmp_path / "e.html").read_text()
    assert "<script>alert(1)</script>" not in html and "<img src=x" not in html
    assert "&lt;script&gt;" in html
    rows = list(csv.reader((tmp_path / "e.csv").open(encoding="utf-8-sig")))
    desc_idx = rows[0].index("rule_description")
    assert rows[1][desc_idx].startswith("'=")
    ws = load_workbook(tmp_path / "e.xlsx")["Findings"]
    assert all(not str(c.value).startswith("=") for row in ws.iter_rows() for c in row if c.value is not None)
    session.close()
