"""Report exporters: HTML, PDF, DOCX, Excel, CSV and JSON."""

from __future__ import annotations

import csv
import json
import logging
from dataclasses import asdict
from pathlib import Path
from typing import Callable

from jinja2 import Environment, FileSystemLoader, select_autoescape

from app.i18n import get_language, num, severity_label, tr
from app.core import paths
from app.core.severity import SEVERITY_COLORS
from app.recommendations.engine import SECTION_TITLES, SECTIONS
from app.reports.data import ReportData
from app.services.session import AnalysisSession
from app.utils.text import csv_safe
from app.utils.timeutil import fmt_ts

log = logging.getLogger(__name__)

def _incident_line(inc) -> str:
    return tr("Assessment: {assessment}. Risk {risk}/100, confidence {confidence}, false-positive probability "
              "{fp}, status {status}.", assessment=tr(inc.assessment), risk=f"{inc.risk_score:.0f}",
              confidence=f"{inc.confidence:.0%}", fp=f"{inc.fp_probability:.0%}", status=tr(inc.status))


_FONT_CACHE: list[tuple[str, str]] = []


def _pdf_fonts() -> tuple[str, str]:
    """Register a TrueType font with Cyrillic support (built-in Helvetica has no Cyrillic glyphs)."""
    if _FONT_CACHE:
        return _FONT_CACHE[0]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (paths.resources_dir() / "fonts" / "DejaVuSans.ttf", paths.resources_dir() / "fonts" / "DejaVuSans-Bold.ttf"),
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    ]
    result = ("Helvetica", "Helvetica-Bold")
    for regular, bold in candidates:
        try:
            if regular.exists():
                pdfmetrics.registerFont(TTFont("WSA-Regular", str(regular)))
                pdfmetrics.registerFont(TTFont("WSA-Bold", str(bold if bold.exists() else regular)))
                pdfmetrics.registerFontFamily("WSA-Regular", normal="WSA-Regular", bold="WSA-Bold",
                                              italic="WSA-Regular", boldItalic="WSA-Bold")
                result = ("WSA-Regular", "WSA-Bold")
                break
        except Exception as exc:  # corrupt font file etc. - fall back to the next candidate
            log.warning("Cannot register PDF font %s: %s", regular, exc)
    _FONT_CACHE.append(result)
    return result


FORMATS = {"pdf": ".pdf", "html": ".html", "docx": ".docx", "xlsx": ".xlsx", "csv": ".csv", "json": ".json"}
_REC_SECTIONS = [(k, SECTION_TITLES[k]) for k in SECTIONS]


# ------------------------------------------------------------------------------------- HTML
def export_html(data: ReportData, path: Path, session: AnalysisSession | None = None) -> Path:
    env = Environment(loader=FileSystemLoader(str(paths.templates_dir())),
                      autoescape=select_autoescape(["html", "j2"]), trim_blocks=True, lstrip_blocks=True)
    template = env.get_template("report.html.j2")
    html = template.render(r=data, ts=fmt_ts, sections=_REC_SECTIONS, t=tr, sl=severity_label, num=num,
                           lang=get_language())
    path.write_text(html, encoding="utf-8")
    return path


# ------------------------------------------------------------------------------------- JSON
def export_json(data: ReportData, path: Path, session: AnalysisSession | None = None) -> Path:
    payload = {
        "report": {"title": data.title, "generated_at": data.generated_at, "generator": data.app},
        "summary": data.summary.to_dict(),
        "incidents": [i.to_dict() for i in data.incidents],
        "alerts": [_group_export(g) for g in data.alerts],
        "hosts": data.hosts,
        "users": data.users,
        "iocs": [i.to_dict() for i in data.iocs],
        "cves": [c.to_dict() for c in data.cves],
        "mitre": data.mitre,
        "recommendations": data.recommendations,
        "timeline": data.timeline,
        "appendix": data.appendix,
    }
    with path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2, default=str)
    return path


def _group_export(g) -> dict:
    d = g.to_dict()
    d.pop("timestamps", None)
    d.pop("representative", None)
    return d


# ------------------------------------------------------------------------------------- CSV
FINDING_COLUMNS = ["id", "severity", "risk_score", "confidence", "fp_probability", "assessment", "title", "rule_id",
                   "rule_level", "rule_description", "category", "agent", "src_ip", "users", "count", "first_seen",
                   "last_seen", "mitre", "cves", "incident_id", "what_happened", "immediate_actions"]
EVENT_COLUMNS = ["id", "timestamp", "severity", "risk_score", "rule_id", "level", "description", "category", "agent",
                 "agent_ip", "src_ip", "dst_ip", "user", "process", "cve", "file_path", "hash", "group_id",
                 "source_file"]


def export_csv(data: ReportData, path: Path, session: AnalysisSession | None = None) -> Path:
    """Writes ``<name>.csv`` (findings) and, when a session is given, ``<name>_events.csv`` (all events)."""
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(FINDING_COLUMNS)
        groups = session.store.iter_all_groups() if session is not None else iter(data.alerts)
        for g in groups:
            writer.writerow([csv_safe(v) for v in (
                g.display_id, g.severity, g.risk_score, g.confidence, g.fp_probability, g.assessment, g.title,
                g.rule_id, g.rule_level, g.rule_description, g.category, g.agent_name, g.src_ip, ";".join(g.users[:20]),
                g.count, fmt_ts(g.first_ts), fmt_ts(g.last_ts), ";".join(g.mitre_ids), ";".join(g.cves),
                g.incident_id, g.what_happened, " | ".join(g.recommendations.get("immediate", [])))])
    if session is not None:
        events_path = path.with_name(path.stem + "_events.csv")
        with events_path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(EVENT_COLUMNS)
            for row in session.store.iter_alert_rows():
                writer.writerow([csv_safe(v) for v in (
                    row["id"], fmt_ts(row["ts"]), row["severity"], row["risk_score"], row["rule_id"], row["level"],
                    row["description"], row["category"], row["agent"], row["agent_ip"], row["src_ip"], row["dst_ip"],
                    row["user"], row["process"], row["cve"], row["file_path"], row["hash"], row["group_id"],
                    row["source_file"])])
    return path


# ------------------------------------------------------------------------------------- Excel
def export_xlsx(data: ReportData, path: Path, session: AnalysisSession | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=False)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2A36")

    def sheet(title: str, headers: list[str], rows, widths: list[int] | None = None):
        ws = wb.create_sheet(tr(title)[:31])
        ws.append([tr(h) for h in headers])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        count = 0
        for row in rows:
            ws.append([csv_safe(v) if isinstance(v, str) else v for v in row])
            count += 1
            if count >= 1_000_000:
                break
        ws.freeze_panes = "A2"
        for idx, width in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        return ws

    ws = wb.active
    ws.title = tr("Summary")
    s = data.summary
    sev_rows = {tr("{severity} events", severity=severity_label(k, upper=False)): k for k in s.severity_counts}
    rows = [("Report", data.title), ("Generated", data.generated_at), ("Mode", tr(s.mode.upper())),
            ("Period", f"{fmt_ts(s.first_ts)} - {fmt_ts(s.last_ts)}"), ("Files", s.files), ("Events", s.events),
            ("Agents", s.agents), ("Rules", s.rules), (tr("Incidents"), s.incidents)]
    rows = [(tr(k), v) for k, v in rows]
    rows += [(label, s.severity_counts[k]) for label, k in sev_rows.items()]
    rows += [(tr(k), v) for k, v in (("Affected hosts", s.affected_hosts), ("Affected users", s.affected_users),
                                     ("External IPs", s.external_ips), ("CVE", s.cves),
                                     ("MITRE techniques", s.mitre_techniques), ("Executive summary", s.executive_summary))]
    for row in rows:
        ws.append([row[0], csv_safe(row[1]) if isinstance(row[1], str) else row[1]])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    ws["B" + str(ws.max_row)].alignment = Alignment(wrap_text=True, vertical="top")
    for label, sev in sev_rows.items():
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            if row[0].value == label:
                row[0].fill = PatternFill("solid", fgColor=SEVERITY_COLORS.get(sev, "#888888").lstrip("#"))

    sheet(tr("Incidents"), ["ID", "Severity", "Risk", "Confidence", "FP probability", "Status", "Assessment", "Title",
                        "Events", "Start", "End", "Hosts", tr("Users"), "Sources", "MITRE", "Chain", "Summary",
                        "Immediate actions"],
          ((i.id, severity_label(i.severity), i.risk_score, i.confidence, i.fp_probability, tr(i.status),
            tr(i.assessment), i.title,
            i.event_count, fmt_ts(i.first_ts), fmt_ts(i.last_ts), ", ".join(i.affected_hosts[:20]),
            ", ".join(i.affected_users[:20]), ", ".join(i.source_ips[:20]),
            ", ".join(m.technique_id for m in i.mitre), " > ".join(i.chain), i.summary,
            "\n".join(i.recommendations.get("immediate", []))) for i in data.incidents),
          [16, 10, 7, 10, 12, 14, 26, 50, 8, 22, 22, 30, 30, 30, 26, 40, 80, 80])
    groups = session.store.iter_all_groups() if session is not None else iter(data.alerts)
    sheet("Findings", FINDING_COLUMNS,
          ((g.display_id, severity_label(g.severity), g.risk_score, g.confidence, g.fp_probability, tr(g.assessment),
            g.title, g.rule_id,
            g.rule_level, g.rule_description, g.category, g.agent_name, g.src_ip, ";".join(g.users[:20]), g.count,
            fmt_ts(g.first_ts), fmt_ts(g.last_ts), ";".join(g.mitre_ids), ";".join(g.cves), g.incident_id,
            g.what_happened, " | ".join(g.recommendations.get("immediate", []))) for g in groups),
          [12, 10, 7, 10, 10, 26, 50, 8, 6, 50, 16, 16, 16, 30, 8, 22, 22, 20, 20, 16, 80, 80])
    sheet("Hosts", ["Host", "Severity", "Max risk", "Events", tr("Incidents"), "Criticality", "IP"],
          ((h["name"], severity_label(h["severity"]), h["max_risk"], h["events"], h["incidents"],
            tr(h["data"].get("criticality", "")),
            h["data"].get("ip", "")) for h in data.hosts), [24, 12, 10, 10, 10, 12, 16])
    sheet(tr("Users"), ["Account", "Severity", "Max risk", "Events", "Hosts"],
          ((u["name"], severity_label(u["severity"]), u["max_risk"], u["events"], ", ".join(u["data"].get("hosts", [])[:10]))
           for u in data.users), [24, 12, 10, 10, 60])
    sheet("IOC", ["Type", "Value", "Verdict", "Occurrences", "Hosts", "Country", "ASN", "Sources"],
          ((i.type, i.value, tr(i.verdict), i.count, ", ".join(i.hosts[:10]), i.country, i.asn,
            "; ".join(f"{s.get('provider')}: {tr(s.get('verdict') or 'unknown')}" for s in i.sources)) for i in data.iocs),
          [10, 50, 12, 12, 30, 10, 10, 60])
    sheet("CVE", ["CVE", "CVSS", "Severity", "Known exploited", "Hosts", "Packages", "Description", "Remediation"],
          ((c.cve, c.cvss, severity_label(c.severity) if c.severity else "",
            tr("YES") if c.known_exploited else (tr("no") if c.kev_checked else tr("unknown (KEV not loaded)")),
            ", ".join(c.hosts[:20]),
            ", ".join(c.packages[:5]), c.description, c.remediation) for c in data.cves),
          [18, 8, 10, 14, 40, 30, 80, 80])
    sheet("MITRE", ["Technique", "Name", "Tactics", "Events", "Confidence", "Evidence"],
          ((m["technique_id"], m["name"], ", ".join(m["tactics"]), m["events"], tr(m["confidence"]),
            "; ".join(m["evidence"])) for m in data.mitre), [12, 36, 30, 10, 12, 80])
    if session is not None:
        sheet("Events", EVENT_COLUMNS,
              ((r["id"], fmt_ts(r["ts"]), severity_label(r["severity"] or ""), r["risk_score"], r["rule_id"], r["level"], r["description"],
                r["category"], r["agent"], r["agent_ip"], r["src_ip"], r["dst_ip"], r["user"], r["process"], r["cve"],
                r["file_path"], r["hash"], r["group_id"], r["source_file"]) for r in session.store.iter_alert_rows()),
              [8, 22, 10, 8, 8, 6, 50, 16, 16, 14, 16, 16, 16, 30, 16, 40, 20, 8, 30])
    wb.save(path)
    return path


# ------------------------------------------------------------------------------------- DOCX
def export_docx(data: ReportData, path: Path, session: AnalysisSession | None = None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    doc.add_heading(data.title, 0)
    meta = doc.add_paragraph(tr("Generated {date} by {app} - Mode: {mode}", date=data.generated_at, app=data.app,
                                mode=tr(data.summary.mode.upper())))
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def table(headers: list[str], rows: list[list], widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = tr(h)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = "" if v is None else str(v)
        return t

    def sev_run(paragraph, severity: str):
        run = paragraph.add_run(f" [{severity_label(severity)}]")
        run.bold = True
        color = SEVERITY_COLORS.get(severity, "#888888").lstrip("#")
        run.font.color.rgb = RGBColor.from_string(color.upper())

    s = data.summary
    if data.has("executive_summary"):
        doc.add_heading(tr("Executive Summary"), 1)
        for para in s.executive_summary.split("\n\n"):
            doc.add_paragraph(para)
    if data.has("statistics"):
        doc.add_heading(tr("Statistics"), 1)
        table(["Metric", "Value"],
              [[tr("{severity} events", severity=severity_label(k, upper=False)), num(v)]
               for k, v in s.severity_counts.items()] +
              [[tr(k), v] for k, v in (("Total events", num(s.events)), ("Files", s.files), ("Agents", s.agents),
                                       ("Rules", s.rules), (tr("Incidents"), s.incidents),
                                       ("Affected hosts", s.affected_hosts), ("Affected users", s.affected_users),
                                       ("External IPs", s.external_ips), ("CVE", s.cves),
                                       ("MITRE techniques", s.mitre_techniques))])
    if data.has("incidents"):
        doc.add_heading(tr("Incidents"), 1)
        for inc in data.incidents[:40]:
            p = doc.add_heading(f"{inc.id} - {inc.title}", 2)
            sev_run(p, inc.severity)
            doc.add_paragraph(_incident_line(inc))
            doc.add_paragraph(inc.summary)
            if inc.chain:
                doc.add_paragraph(tr("Attack chain") + ": " + " -> ".join(inc.chain))
            if inc.mitre:
                doc.add_paragraph("MITRE ATT&CK: " + ", ".join(f"{m.technique_id} {m.name} ({tr(m.confidence)})"
                                                                for m in inc.mitre[:8]))
            for item in inc.evidence[:6]:
                doc.add_paragraph(item, style="List Bullet")
            for key, label in _REC_SECTIONS:
                steps = inc.recommendations.get(key) or []
                if steps:
                    doc.add_paragraph(tr(label)).runs[0].bold = True
                    for step in steps:
                        doc.add_paragraph(step, style="List Number")
    if data.has("critical_alerts") and data.alerts:
        doc.add_heading(tr("Critical and High Alerts"), 1)
        table(["ID", "Severity", "Risk", "Finding", "Host", "Source", "Count"],
              [[g.display_id, severity_label(g.severity), f"{g.risk_score:.0f}", g.title, g.agent_name, g.src_ip,
                num(g.count)]
               for g in data.alerts[:100]])
    if data.has("hosts") and data.hosts:
        doc.add_heading(tr("Affected Hosts"), 1)
        table(["Host", "Severity", "Max risk", "Events", tr("Incidents")],
              [[h["name"], severity_label(h["severity"]), f"{h['max_risk']:.0f}", num(h["events"]), h["incidents"]]
               for h in data.hosts[:60]])
    if data.has("users") and data.users:
        doc.add_heading(tr("Users"), 1)
        table(["Account", "Severity", "Max risk", "Events"],
              [[u["name"], severity_label(u["severity"]), f"{u['max_risk']:.0f}", num(u["events"])]
               for u in data.users[:60]])
    if data.has("iocs") and data.iocs:
        doc.add_heading(tr("Indicators of Compromise"), 1)
        table(["Type", "Value", "Verdict", "Occurrences"],
              [[i.type, i.value, tr(i.verdict), num(i.count)] for i in data.iocs[:100]])
    if data.has("cves") and data.cves:
        doc.add_heading(tr("Vulnerabilities"), 1)
        for c in data.cves[:50]:
            p = doc.add_paragraph()
            run = p.add_run(f"{c.cve} - CVSS {c.cvss if c.cvss is not None else 'n/a'}"
                            + (" - " + tr("KNOWN EXPLOITED") if c.known_exploited else ""))
            run.bold = True
            if c.description:
                doc.add_paragraph(c.description)
            doc.add_paragraph(f"{tr('Hosts')}: {', '.join(c.hosts[:10])}. {tr('Remediation')}: {c.remediation}")
    if data.has("mitre") and data.mitre:
        doc.add_heading("MITRE ATT&CK", 1)
        table(["Technique", "Name", "Tactics", "Events", "Confidence"],
              [[m["technique_id"], m["name"], ", ".join(m["tactics"]), num(m["events"]), tr(m["confidence"])]
               for m in data.mitre[:60]])
    if data.has("recommendations"):
        doc.add_heading(tr("Recommendations"), 1)
        for key, label in _REC_SECTIONS:
            steps = data.recommendations.get(key) or []
            if steps:
                doc.add_heading(tr(label), 2)
                for step in steps:
                    doc.add_paragraph(step, style="List Number")
    if data.has("timeline") and data.timeline:
        doc.add_heading(tr("Timeline"), 1)
        table(["Start", "Severity", "Event"],
              [[fmt_ts(t["ts"]), severity_label(t["severity"]), t["text"]] for t in data.timeline])
    if data.has("appendix"):
        doc.add_heading(tr("Technical Appendix"), 1)
        doc.add_paragraph(tr("Input files") + ": " + ", ".join(data.appendix["files"]))
        doc.add_paragraph(tr("Parsers") + ": " + ", ".join(f"{k} ({num(v)})" for k, v in data.appendix["parsers"].items()))
        doc.add_paragraph(tr("Risk scale") + f": {data.appendix['severity_scale']}.")
        doc.add_paragraph(tr("Statements distinguish observed facts (evidence) from analytical conclusions (assessment)."))
    doc.save(str(path))
    return path


# ------------------------------------------------------------------------------------- PDF
def export_pdf(data: ReportData, path: Path, session: AnalysisSession | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from xml.sax.saxutils import escape

    font, bold_font = _pdf_fonts()
    styles = getSampleStyleSheet()
    for style_name in ("Title", "Heading1", "Heading2", "BodyText", "Normal"):
        styles[style_name].fontName = bold_font if style_name.startswith(("Title", "Heading")) else font
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9, leading=12, alignment=TA_LEFT,
                          fontName=font)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    header = ParagraphStyle("header", parent=small, textColor=colors.white, fontName=bold_font)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, spaceBefore=10, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=8, spaceAfter=4)

    def p(text, style=body):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    def sev_cell(sev: str):
        color = SEVERITY_COLORS.get(sev, "#888888")
        return Paragraph(f'<font color="{color}"><b>{escape(severity_label(sev))}</b></font>', body)

    def tbl(headers, rows, widths):
        data_rows = [[p(tr(h), header) for h in headers]] + rows
        t = Table(data_rows, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A36")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F8")]),
        ]))
        return t

    story = [Paragraph(escape(data.title), styles["Title"]),
             p(tr("Generated {date} by {app} - Mode: {mode}", date=data.generated_at, app=data.app,
                  mode=tr(data.summary.mode.upper())) + f" - {tr('Period')}: {fmt_ts(data.summary.first_ts)} \u2013 "
               f"{fmt_ts(data.summary.last_ts)}", small), Spacer(1, 6)]
    s = data.summary
    if data.has("executive_summary"):
        story.append(Paragraph(tr("Executive Summary"), h1))
        for para in s.executive_summary.split("\n\n"):
            story.append(p(para))
            story.append(Spacer(1, 4))
    if data.has("statistics"):
        story.append(Paragraph(tr("Statistics"), h1))
        rows = [[sev_cell(k), p(num(v))] for k, v in s.severity_counts.items()]
        rows += [[p(tr(k)), p(v)] for k, v in (("Total events", num(s.events)), ("Files", s.files), ("Agents", s.agents),
                                           ("Rules", s.rules), (tr("Incidents"), s.incidents),
                                           ("Affected hosts", s.affected_hosts), ("Affected users", s.affected_users),
                                           ("External IPs", s.external_ips), ("CVE", s.cves),
                                           ("MITRE techniques", s.mitre_techniques))]
        story.append(tbl(["Metric", "Value"], rows, [60, 40]))
    if data.has("incidents") and data.incidents:
        story.append(PageBreak())
        story.append(Paragraph(tr("Incidents"), h1))
        for inc in data.incidents[:40]:
            block = [Paragraph(f"{escape(inc.id)} - {escape(inc.title)} "
                               f'<font color="{SEVERITY_COLORS.get(inc.severity)}">[{escape(severity_label(inc.severity))}]'
                               '</font>', h2),
                     p(_incident_line(inc), small),
                     p(inc.summary)]
            if inc.chain:
                block.append(p(tr("Attack chain") + ": " + " -> ".join(inc.chain)))
            if inc.mitre:
                block.append(p("MITRE ATT&CK: " + ", ".join(f"{m.technique_id} {m.name} ({tr(m.confidence)})"
                                                            for m in inc.mitre[:8]), small))
            for e in inc.evidence[:5]:
                block.append(p("• " + e, small))
            for key, label in _REC_SECTIONS:
                steps = inc.recommendations.get(key) or []
                if steps:
                    block.append(p(tr(label) + ":", small))
                    for n, step in enumerate(steps, start=1):
                        block.append(p(f"  {n}. {step}", small))
            block.append(Spacer(1, 6))
            story.append(KeepTogether(block[:4]))
            story.extend(block[4:])
    if data.has("critical_alerts") and data.alerts:
        story.append(PageBreak())
        story.append(Paragraph(tr("Critical and High Alerts"), h1))
        story.append(tbl(["ID", "Severity", "Risk", "Finding", "Host", "Count"],
                         [[p(g.display_id, small), sev_cell(g.severity), p(f"{g.risk_score:.0f}"), p(g.title, small),
                           p(g.agent_name, small), p(num(g.count))] for g in data.alerts[:120]],
                         [22, 20, 12, 80, 28, 14]))
    if data.has("hosts") and data.hosts:
        story.append(Paragraph(tr("Affected Hosts"), h1))
        story.append(tbl(["Host", "Severity", "Max risk", "Events", tr("Incidents")],
                         [[p(h["name"]), sev_cell(h["severity"]), p(f"{h['max_risk']:.0f}"), p(num(h["events"])),
                           p(h["incidents"])] for h in data.hosts[:60]], [55, 25, 20, 25, 20]))
    if data.has("users") and data.users:
        story.append(Paragraph(tr("Users"), h1))
        story.append(tbl(["Account", "Severity", "Max risk", "Events"],
                         [[p(u["name"]), sev_cell(u["severity"]), p(f"{u['max_risk']:.0f}"), p(num(u["events"]))]
                          for u in data.users[:60]], [60, 30, 25, 25]))
    if data.has("iocs") and data.iocs:
        story.append(Paragraph(tr("Indicators of Compromise"), h1))
        story.append(tbl(["Type", "Value", "Verdict", "Count"],
                         [[p(i.type, small), p(i.value, small), p(tr(i.verdict), small), p(num(i.count), small)]
                          for i in data.iocs[:100]], [18, 110, 25, 20]))
    if data.has("cves") and data.cves:
        story.append(Paragraph(tr("Vulnerabilities"), h1))
        for c in data.cves[:60]:
            story.append(p(f"{c.cve} - CVSS {c.cvss if c.cvss is not None else 'n/a'}"
                           + (" - " + tr("KNOWN EXPLOITED") + " (CISA KEV)" if c.known_exploited else "")
                           + f" - {tr('Hosts')}: {', '.join(c.hosts[:8])}"))
            if c.description:
                story.append(p(c.description, small))
            story.append(p(tr("Remediation") + ": " + c.remediation, small))
            story.append(Spacer(1, 4))
    if data.has("mitre") and data.mitre:
        story.append(Paragraph("MITRE ATT&CK", h1))
        story.append(tbl(["Technique", "Name", "Tactics", "Events", "Confidence"],
                         [[p(m["technique_id"], small), p(m["name"], small), p(", ".join(m["tactics"]), small),
                           p(num(m["events"]), small), p(tr(m["confidence"]), small)] for m in data.mitre[:60]],
                         [22, 50, 55, 18, 22]))
    if data.has("recommendations"):
        story.append(Paragraph(tr("Recommendations"), h1))
        for key, label in _REC_SECTIONS:
            steps = data.recommendations.get(key) or []
            if steps:
                story.append(Paragraph(escape(tr(label)), h2))
                for n, step in enumerate(steps, start=1):
                    story.append(p(f"{n}. {step}"))
    if data.has("timeline") and data.timeline:
        story.append(Paragraph(tr("Timeline"), h1))
        story.append(tbl(["Start", "Severity", "Event"],
                         [[p(fmt_ts(t["ts"]), small), sev_cell(t["severity"]), p(t["text"], small)]
                          for t in data.timeline], [38, 22, 110]))
    if data.has("appendix"):
        story.append(Paragraph(tr("Technical Appendix"), h1))
        story.append(p(tr("Input files") + ": " + ", ".join(data.appendix["files"]), small))
        story.append(p(tr("Parsers") + ": " + ", ".join(f"{k} ({num(v)})" for k, v in data.appendix["parsers"].items()),
                       small))
        story.append(p(tr("Duplicates removed: {duplicates}; malformed records skipped: {errors}; analysis time: "
                          "{seconds} s.", duplicates=data.appendix["duplicates"], errors=data.appendix["parse_errors"],
                          seconds=data.appendix["duration"]), small))
        story.append(p(tr("Risk scale") + ": " + data.appendix["severity_scale"] + ".", small))
        story.append(p(tr("Statements distinguish observed facts (evidence) from analytical conclusions (assessment)."),
                       small))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(15 * mm, 10 * mm, f"{data.app} - {tr('CONFIDENTIAL')}")
        canvas.drawRightString(195 * mm, 10 * mm, tr("Page {page}", page=doc_.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=18 * mm, title=data.title, author=data.app)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


EXPORTERS: dict[str, Callable[[ReportData, Path, AnalysisSession | None], Path]] = {
    "pdf": export_pdf,
    "html": export_html,
    "docx": export_docx,
    "xlsx": export_xlsx,
    "csv": export_csv,
    "json": export_json,
}


def export_report(session: AnalysisSession, out_dir: Path, formats: list[str], base_name: str = "wazuh_report",
                  sections: tuple[str, ...] | None = None, title: str = "Wazuh Security Analysis Report") -> list[Path]:
    from app.reports.data import ALL_SECTIONS, collect_report_data
    from app.services.audit import audit

    out_dir.mkdir(parents=True, exist_ok=True)
    data = collect_report_data(session, tr(title), sections or ALL_SECTIONS)
    written: list[Path] = []
    for fmt in formats:
        fmt = fmt.lower().lstrip(".")
        if fmt == "excel":
            fmt = "xlsx"
        if fmt not in EXPORTERS:
            raise ValueError(f"Unsupported report format: {fmt}")
        path = out_dir / f"{base_name}{FORMATS[fmt]}"
        EXPORTERS[fmt](data, path, session)
        written.append(path)
        log.info("Report written: %s", path)
    audit("report_exported", formats=formats, directory=str(out_dir))
    return written
