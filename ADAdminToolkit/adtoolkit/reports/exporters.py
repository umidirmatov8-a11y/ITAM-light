"""Export of report tables to CSV, XLSX and HTML.

* Every export contains metadata: generation time, domain, domain controller, criteria and data limitations.
* Spreadsheet formula injection is neutralised (cells starting with = + - @ TAB CR are prefixed with an apostrophe in
  CSV and written as plain strings in XLSX).
* Secret attributes are never exported (masked defensively).
"""
from __future__ import annotations

import csv
import html
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from ..security.masking import SECRET_KEYS, mask_text

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class ReportTable:
    title: str
    columns: list[tuple[str, str]]                 # (key, header)
    rows: list[dict]
    criteria: dict[str, Any] = field(default_factory=dict)
    limitations: list[str] = field(default_factory=list)
    generated_at: datetime | None = None
    domain: str = ""
    dc: str = ""
    generated_by: str = ""
    sheet_name: str = ""

    def metadata(self) -> list[tuple[str, str]]:
        gen = (self.generated_at or datetime.now().astimezone()).astimezone()
        meta = [("Отчёт", self.title), ("Дата формирования", gen.strftime("%Y-%m-%d %H:%M:%S %z")),
                ("Домен", self.domain or "—"), ("Контроллер домена", self.dc or "—"),
                ("Сформировал", self.generated_by or "—"), ("Записей", str(len(self.rows)))]
        for k, v in self.criteria.items():
            if v not in (None, "", []):
                meta.append((f"Критерий: {k}", format_value(v)))
        for lim in self.limitations:
            meta.append(("Ограничение данных", lim))
        return meta


def format_value(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Да" if v else "Нет"
    if isinstance(v, datetime):
        return v.astimezone().strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (list, tuple, set)):
        return "; ".join(format_value(x) for x in v)
    if isinstance(v, (bytes, bytearray)):
        return f"<{len(v)} байт>"
    return str(v)


def safe_cell(text: str) -> str:
    """Neutralise spreadsheet formula injection in exported text."""
    if text and text.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def _row_values(table: ReportTable, row: dict) -> list[str]:
    out = []
    for key, _ in table.columns:
        if key.lower() in SECRET_KEYS:
            out.append("********")
            continue
        out.append(mask_text(format_value(row.get(key))))
    return out


def export_csv(table: ReportTable, path: str | Path, *, delimiter: str = ";", include_metadata: bool = True) -> Path:
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)
        if include_metadata:
            for k, v in table.metadata():
                w.writerow([safe_cell(f"# {k}"), safe_cell(v)])
            w.writerow([])
        w.writerow([h for _, h in table.columns])
        for row in table.rows:
            w.writerow([safe_cell(v) for v in _row_values(table, row)])
    return path


def export_xlsx(tables: list[ReportTable] | ReportTable, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tables = [tables] if isinstance(tables, ReportTable) else list(tables)
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    used_names: set[str] = set()

    def sheet_name(name: str) -> str:
        clean = "".join(c for c in name if c not in "[]:*?/\\")[:31] or "Отчёт"
        base, n = clean, 1
        while clean.lower() in used_names:
            n += 1
            clean = f"{base[:28]}_{n}"
        used_names.add(clean.lower())
        return clean

    def put(ws, r, c, value, **style):
        cell = ws.cell(row=r, column=c)
        if isinstance(value, str):
            cell.value = value
            cell.data_type = "s"          # never interpret as a formula
        else:
            cell.value = value
        for k, v in style.items():
            setattr(cell, k, v)
        return cell

    info = wb.create_sheet(sheet_name("Сведения"))
    put(info, 1, 1, "AD Admin Toolkit — сведения об отчёте", font=Font(bold=True, size=13))
    r = 3
    for t in tables:
        put(info, r, 1, t.title, font=Font(bold=True))
        r += 1
        for k, v in t.metadata():
            put(info, r, 1, k)
            put(info, r, 2, v, alignment=Alignment(wrap_text=True, vertical="top"))
            r += 1
        r += 1
    info.column_dimensions["A"].width = 32
    info.column_dimensions["B"].width = 110

    for t in tables:
        ws = wb.create_sheet(sheet_name(t.sheet_name or t.title))
        put(ws, 1, 1, t.title, font=Font(bold=True, size=12))
        gen = (t.generated_at or datetime.now().astimezone()).astimezone().strftime("%Y-%m-%d %H:%M")
        put(ws, 2, 1, f"Сформирован: {gen}; домен: {t.domain or '—'}; DC: {t.dc or '—'}; записей: {len(t.rows)}")
        hdr = 4
        for c, (_, title) in enumerate(t.columns, start=1):
            put(ws, hdr, c, title, fill=header_fill, font=header_font, alignment=Alignment(wrap_text=True, vertical="center"))
        widths = [len(h) for _, h in t.columns]
        for i, row in enumerate(t.rows, start=hdr + 1):
            for c, (key, _) in enumerate(t.columns, start=1):
                raw = row.get(key)
                if key.lower() in SECRET_KEYS:
                    value: Any = "********"
                elif isinstance(raw, bool) or raw is None or isinstance(raw, (list, tuple, set, bytes)):
                    value = format_value(raw)
                elif isinstance(raw, datetime):
                    value = raw.astimezone().replace(tzinfo=None)
                elif isinstance(raw, (int, float)):
                    value = raw
                else:
                    value = mask_text(format_value(raw))
                cell = put(ws, i, c, value)
                if isinstance(value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                widths[c - 1] = max(widths[c - 1], min(60, len(format_value(value))))
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = max(10, min(62, w + 2))
        if t.columns:
            ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(t.columns))}{hdr + max(1, len(t.rows))}"
            ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
        if t.limitations:
            lr = hdr + len(t.rows) + 2
            put(ws, lr, 1, "Ограничения данных:", font=Font(bold=True))
            for j, lim in enumerate(t.limitations, start=1):
                put(ws, lr + j, 1, lim)
    wb.save(path)
    return path


_HTML_STYLE = """
body{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#1b2430;background:#fff}
h1{font-size:22px;margin:0 0 4px} h2{font-size:17px;margin:28px 0 8px;border-bottom:2px solid #1f3a5f;padding-bottom:4px}
.meta{color:#556;font-size:12px;margin-bottom:12px} table{border-collapse:collapse;width:100%;font-size:12px;margin:6px 0 10px}
th{background:#1f3a5f;color:#fff;text-align:left;padding:5px 6px} td{border-bottom:1px solid #dde3ea;padding:4px 6px;vertical-align:top}
tr:nth-child(even) td{background:#f5f7fa} .lim{font-size:12px;color:#7a4b00;background:#fff6e0;border-left:4px solid #f0a000;padding:6px 10px;margin:6px 0}
.kpis{display:flex;flex-wrap:wrap;gap:10px;margin:10px 0} .kpi{border:1px solid #d5dce6;border-radius:6px;padding:8px 12px;min-width:150px}
.kpi b{display:block;font-size:20px}
.sev-critical,.sev-high{color:#b3261e;font-weight:600}.sev-medium{color:#a15c00}.sev-low{color:#215f9a}
"""


def export_html(tables: list[ReportTable], path: str | Path, *, title: str, kpis: list[tuple[str, str]] | None = None,
                intro: str = "") -> Path:
    path = Path(path)
    e = html.escape
    parts = [f"<!doctype html><html lang='ru'><head><meta charset='utf-8'><title>{e(title)}</title>"
             f"<style>{_HTML_STYLE}</style></head><body><h1>{e(title)}</h1>"]
    if tables:
        t0 = tables[0]
        gen = (t0.generated_at or datetime.now().astimezone()).astimezone().strftime("%Y-%m-%d %H:%M")
        parts.append(f"<div class='meta'>Сформирован: {e(gen)} · Домен: {e(t0.domain or '—')} · "
                     f"Контроллер: {e(t0.dc or '—')} · Сформировал: {e(t0.generated_by or '—')}</div>")
    if intro:
        parts.append(f"<p>{e(intro)}</p>")
    if kpis:
        parts.append("<div class='kpis'>" + "".join(f"<div class='kpi'>{e(k)}<b>{e(v)}</b></div>" for k, v in kpis) + "</div>")
    for t in tables:
        parts.append(f"<h2>{e(t.title)}</h2>")
        crit = [f"{e(k)}: {e(format_value(v))}" for k, v in t.criteria.items() if v not in (None, "", [])]
        if crit:
            parts.append("<div class='meta'>Критерии: " + "; ".join(crit) + "</div>")
        for lim in t.limitations:
            parts.append(f"<div class='lim'>{e(lim)}</div>")
        parts.append("<table><tr>" + "".join(f"<th>{e(h)}</th>" for _, h in t.columns) + "</tr>")
        for row in t.rows:
            cells = []
            for (key, _), v in zip(t.columns, _row_values(t, row)):
                cls = f" class='sev-{e(str(row.get('severity_key', '')))}'" if key == "severity" else ""
                cells.append(f"<td{cls}>{e(v)}</td>")
            parts.append("<tr>" + "".join(cells) + "</tr>")
        parts.append("</table>")
        if not t.rows:
            parts.append("<div class='meta'>Нет записей.</div>")
    parts.append("<div class='meta'>Отчёт сформирован AD Admin Toolkit по данным LDAP. Журнал приложения не является "
                 "полным аудитом Active Directory.</div></body></html>")
    path.write_text("".join(parts), encoding="utf-8")
    return path


def export_table(table: ReportTable, path: str | Path) -> Path:
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return export_xlsx(table, path)
    if suffix in (".html", ".htm"):
        return export_html([table], path, title=table.title)
    return export_csv(table, path)
