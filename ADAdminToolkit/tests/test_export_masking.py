import csv
import logging

import pytest
from openpyxl import load_workbook

from adtoolkit.core.errors import AuthenticationError, describe_exception
from adtoolkit.reports.exporters import ReportTable, export_csv, export_html, export_xlsx, safe_cell
from adtoolkit.security.audit_log import OperationJournal
from adtoolkit.security.masking import MASK, SecretMaskingFilter, mask_mapping, mask_text


def table():
    return ReportTable("Тест", [("name", "Имя"), ("formula", "Формула"), ("password", "Пароль"), ("when", "Когда")],
                       [{"name": "Иванов", "formula": "=HYPERLINK(\"http://evil\")", "password": "P@ssw0rd!", "when": None},
                        {"name": "+cmd|' /C calc'!A0", "formula": "-2+3", "password": "x", "when": True}],
                       criteria={"Период": "90 дн."}, limitations=["lastLogonTimestamp с задержкой"], domain="demo.local",
                       dc="dc01.demo.local", generated_by="DEMO\\admin")


def test_csv_export_metadata_and_injection(tmp_path):
    p = export_csv(table(), tmp_path / "r.csv")
    raw = p.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert "# Домен;demo.local" in text and "# Контроллер домена;dc01.demo.local" in text
    assert "# Критерий: Период;90 дн." in text and "Ограничение данных" in text
    rows = list(csv.reader(text.splitlines(), delimiter=";"))
    data = rows[rows.index(["Имя", "Формула", "Пароль", "Когда"]) + 1:]
    assert data[0][1].startswith("'=") and data[1][0].startswith("'+") and data[1][1] == "'-2+3"
    assert "P@ssw0rd!" not in text and data[0][2] == "********"


def test_xlsx_export(tmp_path):
    p = export_xlsx([table(), table()], tmp_path / "r.xlsx")
    wb = load_workbook(p)
    assert wb.sheetnames[0] == "Сведения" and len(wb.sheetnames) == 3
    ws = wb[wb.sheetnames[1]]
    assert ws.cell(4, 1).value == "Имя"
    cell = ws.cell(5, 2)
    assert cell.value == "=HYPERLINK(\"http://evil\")" and cell.data_type == "s"     # stored as text, not a formula
    assert ws.cell(5, 3).value == "********"
    info = wb["Сведения"]
    values = [c.value for row in info.iter_rows() for c in row if c.value]
    assert "demo.local" in values and any("lastLogonTimestamp" in str(v) for v in values)


def test_html_export_escapes(tmp_path):
    t = table()
    t.rows[0]["name"] = "<script>alert(1)</script>"
    p = export_html([t], tmp_path / "r.html", title="Отчёт <b>")
    html = p.read_text(encoding="utf-8")
    assert "<script>" not in html and "&lt;script&gt;" in html and "P@ssw0rd!" not in html


def test_safe_cell():
    assert safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell("plain") == "plain"


@pytest.mark.parametrize("text", [
    "bind failed: password=Secr3t! user=admin",
    '{"password": "Secr3t!", "user": "admin"}',
    "unicodePwd: Secr3t!",
    "net use /p:Secr3t!",
])
def test_mask_text(text):
    out = mask_text(text)
    assert "Secr3t!" not in out and MASK in out


def test_mask_literal_and_mapping():
    assert "hunter2" not in mask_text("value hunter2 here", ["hunter2"])
    m = mask_mapping({"new_password": "x", "nested": {"unicodePwd": b"abc", "ok": "y"}, "list": [{"secret": 1}]})
    assert m["new_password"] == MASK and m["nested"]["unicodePwd"] == MASK and m["nested"]["ok"] == "y"
    assert m["list"][0]["secret"] == MASK


def test_logging_filter_masks(caplog):
    logger = logging.getLogger("test.mask")
    caplog.set_level(logging.INFO)
    handler_filter = SecretMaskingFilter()
    logger.addFilter(handler_filter)
    logger.info("connecting with password=%s", "TopSecret1")
    assert "TopSecret1" not in caplog.text and MASK in caplog.text


def test_errors_never_show_secrets():
    e = AuthenticationError("failed for password=abc123", details="credentials: abc123")
    assert "abc123" not in e.full_text() and "abc123" not in describe_exception(e)


def test_journal_and_database_never_store_secrets(db):
    j = OperationJournal(db)
    j.record("user.reset_password", "CN=x", "success", details={"password": "Sup3r!", "must_change": True},
             error="password=Sup3r!")
    row = db.operations()[0]
    assert "Sup3r!" not in str(row)
    with pytest.raises(ValueError):
        db.set_setting("ldap_password", "x")
    db.save_profile({"name": "p", "server": "dc", "password": "x"})
    assert "password" not in db.list_profiles()[0]
    db.save_template("user", "t", {"ou_dn": "OU=x", "password": "zzz"})
    assert db.templates("user")[0]["password"] == MASK
